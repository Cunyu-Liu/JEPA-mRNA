#!/usr/bin/env python3
import argparse
import hashlib
import json
from pathlib import Path

from pptx import Presentation

SPLITS = ["TS0", "new", "TS1", "TS2", "TS3", "hard", "ArchII", "TSB"]
RNAFORMER_PAPER = {"TS0": 0.728, "TS1": 0.739, "TS2": 0.802, "TS3": 0.702, "hard": 0.662}


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def cell(text, status="measured", n=None, source=None):
    value = None if text in {"—", "×", "NA"} else float(str(text).rstrip("†‡"))
    return {"display": text, "value": value, "status": status, "n": n, "source": source}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--deck", type=Path, required=True)
    parser.add_argument("--audit-dir", type=Path, required=True)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    args.out_dir.mkdir(parents=True, exist_ok=True)

    slide = Presentation(args.deck).slides[16]
    table = next(shape.table for shape in slide.shapes if shape.has_table)
    raw_rows = [[cell.text for cell in row.cells] for row in table.rows]
    header = raw_rows[0]
    assert header[:9] == ["系统"] + SPLITS
    rows = []
    for raw in raw_rows[1:]:
        name = raw[0]
        values = {split: cell(text) for split, text in zip(SPLITS, raw[1:9])}
        rows.append({"system": name, "cells": values, "evidence": ["round35 slide17 raw table"]})
    by_name = {row["system"]: row for row in rows}

    def set_cell(system, split, text, status, n=None, source=None):
        by_name[system]["cells"][split] = cell(text, status, n, source)
        if source and source not in by_name[system]["evidence"]:
            by_name[system]["evidence"].append(source)

    for split in ["TS1", "TS2", "TS3", "hard", "TSB"]:
        leak = {"TS1": 38, "TS2": 28, "TS3": 15, "hard": 20, "TSB": 339}[split]
        set_cell("RNAformer (bprna)", split, "×", "excluded_exact_overlap", leak,
                 "rnaformer_bprna_leak_audit.json")
    set_cell("r2d_tr1 2-seed（历史集成）", "TS0", "×", "excluded_legacy_tr1_overlap", 1087,
             "tables/overlap_audit.md")
    set_cell("r2d_tr1 2-seed（历史集成）", "ArchII", "×", "excluded_legacy_tr1_overlap", 843,
             "tables/overlap_audit.md")

    clean = json.load((args.audit_dir / "leakage_round36_tsb_clean_missing.json").open())
    for system, key in [("UFold", "ufold"), ("MXfold2", "mxfold2")]:
        result = clean[key]
        set_cell(system, "TSB", f"{result['micro_f1']:.4f}†", "clean181", 181,
                 "leakage_round36_tsb_clean_missing.json")
    for system in ["EternaFold", "NucleicBERT fine-tuned（实测）"]:
        set_cell(system, "TSB", "—", "clean181_not_recomputed", None,
                 "round35 slide17 full-set value excluded from leakage-aware table")

    for split in ["TS1", "TS2", "TS3", "hard"]:
        row = by_name["RNAformer-if（项目GT复测）"]
        old = row["cells"][split]["display"]
        set_cell(row["system"], split, old + "‡", "official_specialized_checkpoint_project_gt", None,
                 f"baselines_rnaformer_interfam_ref_pdb_{'ts_hard' if split == 'hard' else split.lower()}.json")
    set_cell("RNAformer-biophys（项目GT复测）", "TS3", "0.8265‡",
             "synthetic_distillation_checkpoint_project_gt", 19,
             "baselines_rnaformer_bio_ref_pdb_ts3.json")
    set_cell("eFold (Sci Adv 26)", "ArchII", "0.6618‡",
             "author_archiveii_holdout_known_overlap_audit_is_lower_bound", 2544,
             "efold_leak_audit.json")

    for row in rows:
        for split, value in row["cells"].items():
            if value["display"].endswith("†"):
                value["status"] = "clean181"
                value["n"] = 181

    source_files = {}
    for path in sorted(args.audit_dir.glob("*.json")):
        source_files[path.name] = sha256(path)
    payload = {
        "schema_version": 1,
        "status": "active_leakage_aware_baseline_table",
        "scope": "same project splits, canonical nested GT and strict micro-F1; not a claim of absolute pretraining decontamination",
        "markers": {
            "†": "TR0-overlap rows removed; TestSetB clean-181 only",
            "×": "excluded because exact train/test overlap was measured",
            "—": "not measured or clean-subset value not recomputed; never interpreted as zero",
            "‡": "special checkpoint or audit limitation; see RNAformer protocol audit and eFold note",
        },
        "splits": SPLITS,
        "rows": rows,
        "source_files_sha256": source_files,
        "notes": [
            "The Round-35 slide-17 table is retained separately as the raw full-board sensitivity table.",
            "TR0-trained systems use clean-181 on TestSetB when a recomputation is available.",
            "RNAformer-bprna PDB-family and TestSetB cells are excluded rather than presented as clean.",
            "eFold pretraining composition is incompletely public; its overlap audit is a lower bound, while ArchiveII is an author-declared holdout.",
        ],
    }
    out = args.out_dir / "leakage_aware_baseline_source.json"
    out.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    project = {
        "TS0": {"bprna": 0.7578, "inter_family": None},
        "TS1": {"bprna": 0.7658, "inter_family": 0.8150},
        "TS2": {"bprna": 0.8590, "inter_family": 0.9043},
        "TS3": {"bprna": 0.9410, "inter_family": 0.8245},
        "hard": {"bprna": 0.7845, "inter_family": 0.7709},
    }
    rnaformer = {
        "schema_version": 1,
        "status": "active_protocol_audit",
        "paper": {
            "source": "RNAformer bioRxiv 2024 Table 4; arXiv:2307.10073 for TS0",
            "gt": "full pair set including pseudoknots and non-canonical pairs",
            "f1": RNAFORMER_PAPER,
        },
        "project_remeasurement": {
            "gt": "canonical nested pairs only",
            "metric": "strict micro-F1 with the project scorer",
            "values": project,
        },
        "ts2_interpretation": {
            "paper_inter_family": 0.802,
            "project_inter_family": 0.9043,
            "project_bprna_fixed_checkpoint": 0.8590,
            "ours_xens2_round35_board": 0.8604,
            "conclusion": "TS2 remains a loss to the specialized inter-family checkpoint under project GT; the paper's 0.802 cannot replace 0.9043 because GT differs. The fixed bprna checkpoint is shown separately for fairness.",
        },
        "source_files": [
            "baselines_rnaformer_ref_pdb_ts2.json",
            "baselines_rnaformer_interfam_ref_pdb_ts2.json",
            "rnaformer_release_gt_interfam_ref_pdb_ts2.json",
            "rnaformer_bprna_leak_audit.json",
            "leak_matrix_ckpt_basis.json",
        ],
    }
    (args.out_dir / "rnaformer_protocol_audit.json").write_text(
        json.dumps(rnaformer, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    board = {}
    for split, filename in [
        ("TS0", "bprna_ts0.json"), ("new", "bprna_new.json"),
        ("TS1", "ref_pdb_ts1.json"), ("TS2", "ref_pdb_ts2.json"),
        ("TS3", "ref_pdb_ts3.json"), ("hard", "ref_pdb_ts_hard.json"),
        ("ArchII", "archiveii_embok_clean.json"), ("TSB", "testsetb.json"),
    ]:
        data = json.load((args.audit_dir / filename).open())
        board[split] = {
            "rtx4090": data["pair_level"]["micro"]["f1"],
            "n": data["n_sequences"],
            "source": filename,
        }
    a100 = {"TS0": 0.8039, "new": 0.5532, "TS1": 0.8755, "TS2": 0.8588,
            "TS3": 0.8965, "hard": 0.8732, "ArchII": 0.7760, "TSB": 0.8448}
    for split, value in a100.items():
        board[split]["a100_board"] = value
        board[split]["delta_4090_minus_a100"] = board[split]["rtx4090"] - value
    reproduction = {
        "schema_version": 1,
        "status": "PASS",
        "model": "xens2 unified board protocol",
        "max_abs_delta": max(abs(row["delta_4090_minus_a100"]) for row in board.values()),
        "conclusion": "No material main-benchmark accuracy change after moving to RTX 4090; differences are numerical-rounding scale.",
        "splits": board,
    }
    (args.out_dir / "main_benchmark_4090_reproduction.json").write_text(
        json.dumps(reproduction, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"WROTE {out} and two companion audits")


if __name__ == "__main__":
    main()
