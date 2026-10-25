#!/usr/bin/env python3
"""RNA-JEV 正式 PPT 的内容、字体、行距、重复图和秘密信息专项检查。"""
from pathlib import Path
import argparse
import hashlib
import re
from collections import Counter

from pptx import Presentation


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--draft", action="store_true")
    args = parser.parse_args()
    prs = Presentation(args.pptx)
    assert len(prs.slides) == 19, len(prs.slides)
    fonts, sizes, spacing_errors, bounds_errors = Counter(), [], [], []
    texts, images = [], []
    for slide_index, slide in enumerate(prs.slides, 1):
        assert len(slide.shapes) > 0, f"blank slide {slide_index}"
        slide_text = []
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left + shape.width > prs.slide_width + 2 or shape.top + shape.height > prs.slide_height + 2:
                bounds_errors.append((slide_index, shape.name))
            if shape.shape_type == 13:
                digest = hashlib.sha256(shape.image.blob).hexdigest()
                images.append((slide_index, digest))
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    if paragraph.text.strip():
                        slide_text.append(paragraph.text)
                    for run in paragraph.runs:
                        if not run.text:
                            continue
                        if run.font.name:
                            fonts[run.font.name] += 1
                        if run.font.size:
                            size = run.font.size.pt
                            sizes.append(size)
                            if paragraph.line_spacing is not None:
                                expected = size * 1.5
                                actual = paragraph.line_spacing.pt
                                if abs(actual - expected) > 0.03:
                                    spacing_errors.append((slide_index, run.text[:24], size, actual, expected))
        texts.append("\n".join(slide_text))
    assert set(fonts).issubset({"微软雅黑", "Microsoft YaHei"}), fonts
    assert not spacing_errors, spacing_errors[:20]
    assert not bounds_errors, bounds_errors[:20]
    assert min(sizes) >= 7.0, min(sizes)
    all_text = "\n".join(texts)
    forbidden = ["lzr9", "28055 root@", "BEGIN OPENSSH PRIVATE KEY", "password", "lorem ipsum"]
    for token in forbidden:
        assert token.lower() not in all_text.lower(), token
    if not args.draft:
        for token in ["DRAFT", "等待完整实测矩阵", "缺一格即不出主图", "TODO", "TBD"]:
            assert token not in all_text, token
        assert len(images) == 9, f"expected nine unique main raster figures, got {len(images)}"
        assert len(set(digest for _, digest in images)) == len(images), "duplicate raster image across slides"
        assert "统一模型口径：单模型 4/8，xens2 为 6/8" in texts[3]
        assert "RNAformer-interfam" in texts[2] and "TSB 除外" in texts[2]
        assert "主判定参照（TSB 除外）" in texts[2]
        assert "A100 上是 RNAformer 与 UFold" in texts[7]
        assert "r2dtr1c-s0" in texts[4] and "0.667971" in texts[4]
        assert "0.5532" in texts[14] and "0.6132" in texts[14]
    print(f"DECK_QA_PASS slides={len(prs.slides)} images={len(images)} unique_images={len(set(d for _, d in images))}")
    print("FONTS", dict(fonts))
    print("FONT_SIZE_RANGE", min(sizes), max(sizes))
    print("LINE_SPACING", "1.5x verified per explicit run")


if __name__ == "__main__":
    main()
