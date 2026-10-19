#!/usr/bin/env python3
"""Surgical data refresh on slides 14 & 21 (tables only, style untouched):

slide 14 (SOTA main table):
- TS2 xens2 0.8588 -> 0.8604 (v3.20..v3.25 final verdict, 7-front)
- TS2 vs-ref "- 0.046" -> "- 0.044"

slide 21 (speed table):
- row 4 (UFold reference): judgment cell gains the final state
- update row 3 judgment: v2 48.6x was the 10-16 state; append v4 outcome
  in the same cell; update the last text box (TB below table) to the
  10-19 final state (124.78, 8.6x same-node) — text content only,
  geometry/typography untouched.

All numbers pre-checked: ledger 15.79/15.81/15.83, checkers, PDF 8/8."""
from pptx import Presentation
import copy

P = "/Users/liucunyu/Documents/课题idea 数据 结论/浙江大学/RNA 相关模型/RNA-JEV/ppt/RNA-JEV_项目进展_20260924.pptx"
p = Presentation(P)

def set_cell_text(cell, new_text):
    """Replace text preserving the first run's formatting."""
    para = cell.text_frame.paragraphs[0]
    runs = para.runs
    if runs:
        # keep first run with new text, drop others
        runs[0].text = new_text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        para.add_run().text = new_text  # inherits para/level defaults

# --- slide 14: TS2 row ---
s14 = p.slides[13]
t14 = [sh for sh in s14.shapes if sh.has_table][0].table
rows14 = list(t14.rows)
# find TS2 row
for r in rows14:
    cells = list(r.cells)
    if cells[0].text.strip() == "TS2":
        assert "0.8588" in cells[2].text, cells[2].text
        set_cell_text(cells[2], "0.8604")
        assert "0.046" in cells[4].text, cells[4].text
        set_cell_text(cells[4], "− 0.044")
        print("slide 14 TS2 row updated")
        break

# --- slide 21: speed table + bottom text ---
s21 = p.slides[20]
t21 = [sh for sh in s21.shapes if sh.has_table][0].table
rows21 = list(t21.rows)
# row 3 (C vectorized v2) judgment: append v4 state
c = list(rows21[3].cells)[4]
old = c.text.strip()
new = "48.6× → v4 批轴+fp16：B=256 达 124.78（8.6× 同节点 RiNALMo-ft）"
set_cell_text(c, new)
print("slide 21 row4 judgment:", old[:30], "->", new[:40])
# row 4 (UFold) judgment: close the gap statement
c = list(rows21[4].cells)[4]
old = c.text.strip()
set_cell_text(c, "已被 v4 超过：124.78 vs 56.49（同口径对比 2.2×，批口径 8.6× 同节点参照）")
print("slide 21 row5 judgment updated")

# bottom text box (TextBox 930-family, contains 48.6 narrative)
for sh in s21.shapes:
    if sh.has_text_frame:
        t = sh.text_frame.text
        if "下一目标：B=64" in t and "48.6" in t:
            paras = sh.text_frame.paragraphs
            runs = paras[0].runs
            new_text = ("速度工程线收官（10-19）：批化 + C 向量化解码 + fp16 批轴 = 124.78 seq/s"
                        "（B=256，F1 逐位无损 0.668），同节点 RiNALMo-ft 复现 14.5 seq/s → 8.6×；"
                        "对 UFold 同口径 56.49 = 2.2×。数据：speed_benchmark.json + §15.79/15.81。")
            if runs:
                runs[0].text = new_text
                for r in runs[1:]:
                    r._r.getparent().remove(r._r)
            print("slide 21 bottom narrative updated")
# headline subtitle (TextBox 2) mentions "距离 SOTA 单模速度 2.1x"
for sh in s21.shapes:
    if sh.has_text_frame and "2.1×" in sh.text_frame.text and sh.name in ("TextBox 2",):
        runs = sh.text_frame.paragraphs[0].runs
        old = sh.text_frame.text
        new_text = old.replace("48.6× 无损加速，距离 SOTA 单模速度 2.1×", "124.78 seq/s 收官（8.6× 同节点 RiNALMo-ft）")
        if new_text != old and runs:
            runs[0].text = new_text
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            print("slide 21 headline updated")

p.save(P)
print("saved")
