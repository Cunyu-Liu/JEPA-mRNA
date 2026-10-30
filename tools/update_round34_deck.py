#!/usr/bin/env python3
"""Apply verified Round-33 figure replacements and rebuild Fig. 6 as native vectors."""
from pathlib import Path
from copy import deepcopy
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement

SRC = Path("/private/tmp/rnajev_audit/RNA-JEV_论文逻辑重构_20261026_导师讨论_preview_draft.pptx")
OUT = Path(os.environ.get("DECK_OUT", "/Users/liucunyu/Documents/all_code/ZJU/mRNA_editflow/temp_20261030_02/RNA-JEV_论文逻辑重构_20261030_导师讨论_round34_working.pptx"))
FIG = Path("/Users/liucunyu/Documents/all_code/ZJU/mRNA_editflow/temp_20261028_03/figure_assets")
FONT = "微软雅黑"
C = {
    "ink": "17243A", "muted": "667085", "faint": "98A2B3", "blue": "0F4D92",
    "blue2": "3775BA", "teal": "0E7C86", "green": "287A5B", "red": "B64342",
    "gold": "C98A22", "surface": "F5F7FA", "rule": "DDE3EA", "white": "FFFFFF",
    "paleBlue": "EAF1F8", "paleTeal": "E9F5F4", "paleGold": "FBF4E4", "dark": "101B31",
}


def rgb(key):
    return RGBColor.from_string(C.get(key, key))


def remove(shape):
    shape._element.getparent().remove(shape._element)


def text(slide, value, x, y, w, h, size=10, color="ink", bold=False,
         align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.02):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = Inches(margin)
    frame.margin_top = frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    frame.word_wrap = True
    lines = str(value).split("\n")
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.line_spacing = Pt(size * 1.5)
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = rgb(color)
    return shape


def box(slide, x, y, w, h, fill="white", line="rule", radius=True, width=0.7):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line)
    shape.line.width = Pt(width)
    return shape


def labeled_box(slide, label, x, y, w, h, fill="white", line="rule", size=9,
                color="ink", bold=False, align=PP_ALIGN.CENTER):
    box(slide, x, y, w, h, fill, line)
    return text(slide, label, x + 0.06, y + 0.03, w - 0.12, h - 0.06, size, color, bold, align)


def connector(slide, x1, y1, x2, y2, color="muted", width=1.1, arrow=True):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    if arrow:
        ln = shape._element.spPr.get_or_add_ln()
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        ln.append(tail)
    return shape


def matrix(slide, x, y, cell=0.16, n=6, palette=("paleBlue", "blue2", "teal")):
    for i in range(n):
        for j in range(n):
            key = palette[0] if i == j else palette[1] if (i + j) % 3 == 0 else palette[2] if i < j else "surface"
            box(slide, x + j * cell, y + i * cell, cell - 0.025, cell - 0.025,
                key, "white", radius=False, width=0.25)


def replace_picture(slide, image_path):
    pictures = [shape for shape in slide.shapes if shape.shape_type == 13]
    assert len(pictures) == 1, len(pictures)
    image_part, r_id = slide.part.get_or_add_image_part(str(image_path))
    pictures[0]._element.blipFill.blip.rEmbed = r_id


def text_frames(container):
    for shape in container.shapes:
        if getattr(shape, "has_text_frame", False):
            yield shape.text_frame
        elif getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def replace_run_text(container, old, new, expected):
    count = 0
    for frame in text_frames(container):
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    count += 1
    assert count == expected, (old, count, expected)


def update_secondary_data_status(prs):
    replace_picture(prs.slides[10], FIG / "fig07_nucleicbert_combined_working_3840x1920.png")
    replace_run_text(
        prs.slides[10],
        "8/8 仅指主 benchmark；secondary structure/contact-distance/shuffle 缺官方资产，保持 blocked，不作全覆盖声明。",
        "8/8 仅指主 benchmark；secondary-structure 数据门已过、五种子待验收；contact-distance/shuffle 保持 blocked，不作全覆盖声明。",
        1,
    )
    replace_run_text(
        prs.slides[14],
        "NB 其余官方资产 blocked；MRL 已验收为负，TE/EL 十折未闭环",
        "NB 二级结构待五种子；contact/distance/shuffle blocked；MRL 为负，TE/EL 未闭环",
        1,
    )
    replace_run_text(
        prs.slides[10].notes_slide,
        "8/8 仅指主 benchmark；secondary structure/contact-distance/shuffle 缺官方资产，保持 blocked，不作全覆盖声明。",
        "8/8 仅指主 benchmark；secondary-structure 数据门已过、五种子待验收；contact-distance/shuffle 保持 blocked，不作全覆盖声明。",
        1,
    )


def update_version_labels(prs):
    replace_run_text(prs.slides[0], "2026-10-26 · Round 32", "2026-10-30 · Round 34", 1)
    replace_run_text(prs.slides[14], "结论口径：Round 32 统一复核", "结论口径：Round 34 统一复核", 1)
    replace_run_text(prs.slides[20], "RNA-JEV / ROUND 32", "RNA-JEV / ROUND 34", 1)


def update_speed_provenance(prs):
    replacements = {
        5: [
            ("主图 2｜r2d 单模型 124.78 seq/s，F1 逐位不变",
             "主图 2｜历史 head-only 峰值 124.78 seq/s，禁止跨口径比较", 2),
            ("RTX 4090、TS0、端到端含 frozen encoder + 9.3M head + C decode。",
             "RTX 4090、TS0；JEV 梯图为 cached-embedding/head-only，812,375 head；非 raw-sequence 端到端。", 1),
            ("124.78 属于 r2dtr1c-s0（F1=0.667971），不与 xens2 的 0.8039 拼接；A100 与 forward-only 参照分别标注。",
             "124.78 属于 r2dtr1c-s0 的历史 head-only 峰值；当前 raw-sequence B=1 为 22.28 seq/s（F1=0.668181），二者不得混用。", 2),
        ],
        9: [
            ("主图 5｜真实帕累托集：A100 上是 RNAformer 与 UFold",
             "主图 5｜历史 A100 帕累托集；4090 共同协议重测中", 2),
            ("同节点、同 split、strict micro-F1；r2d 在 A100 被 RNAformer 同时以更高 F1 和吞吐支配。",
             "当前图仅作 A100 历史参照；正式版使用同一 4090、TS0、raw-sequence 端到端 B=1/2/4 结果。", 1),
            ("4090 的 r2d=124.78 seq/s 仅作硬件更新，不混入 A100 前沿；不虚构 xens2 的吞吐或联合最优点。",
             "旧 124.78 为 cached-embedding/head-only，不能混入 raw-sequence 前沿；当前仅四个系统 B=1 已通过，完整网格待验收。", 2),
        ],
        15: [
            ("r2d 在 4090 达 124.78 seq/s，批大小改变不改预测",
             "4090 raw-seq B=1：r2d 22.28 seq/s；旧 head-only 峰值 124.78", 1),
            ("A100 r2d 不在 Pareto 前沿；不把该速度与 xens2 F1 拼成联合点",
             "共同 batch 网格未闭环；禁止跨口径拼接速度或 xens2 F1", 1),
        ],
        17: [
            ("附录 B / 速度协议", "附录 B / 历史 head-only 速度协议", 1),
            ("速度测试串行执行；不把跨硬件或 forward-only 读数写成同口径。",
             "本页为历史 head-only 阶梯；不得与 raw-sequence 或 forward-only 读数混用。", 1),
            ("配置", "历史 head-only 配置", 1),
            ("终值；峰值", "历史峰值", 1),
            ("124.78 属于 r2dtr1c-s0，不是 xens2；F1=0.667971。",
             "124.78 属于 r2dtr1c-s0 的历史 head-only 峰值，不是 xens2；F1=0.667971。", 1),
            ("端到端范围", "计时范围", 1),
            ("含 frozen RiNALMo 前向、9.3M 头与 C 解码；不是 encoder-cache 口径。",
             "含 cached embedding、812,375 head 与 C 解码；不包含 raw-sequence 编码器前向。", 1),
        ],
        20: [
            ("r2d 在 4090 达 124.78 seq/s",
             "raw-seq B=1：r2d 22.28 seq/s", 1),
            ("A100 前沿为 RNAformer/UFold", "旧 head-only 峰值 124.78；共同 batch 网格待验收", 1),
        ],
    }
    for slide_number, items in replacements.items():
        slide = prs.slides[slide_number - 1]
        for old, new, expected in items:
            replace_run_text(slide, old, new, expected - 1 if expected == 2 else expected)
            if expected == 2:
                replace_run_text(slide.notes_slide, old, new, 1)


def update_downstream_provenance(prs):
    replacements = {
        11: [
            ("主图 7｜Jev-on-NucleicBERT：8/8 split 均正增益",
             "主图 7｜主 benchmark 8/8 正迁移；原论文三任务不显著为正", 2),
            ("分组柱状图：Δ +0.055 至 +0.234，平均 +0.134；OOD/家族 split 增益更大。",
             "主 benchmark 与原论文任务分开展示；acceptor/donor/CPEB3 全量五种子均不显著为正。", 1),
            ("回答兼容性问题：同一 Jev 决策头不依赖 RiNALMo 特有表征；ArchII 使用匹配的 trunc598 口径。",
             "8/8 仅指主 benchmark；secondary structure/contact-distance/shuffle 缺官方资产，保持 blocked，不作全覆盖声明。", 2),
            ("ROUND26 / T-A99；NB-ft 与 Jev-on-NB 8-split",
             "ROUND26/T-A99 + round34 downstream QA；其余原论文任务待官方资产", 2),
        ],
        12: [
            ("主图 8｜下游头迁移：ncRNA 三指标提升，回归任务小幅为正",
             "主图 8A｜官方分类未显著为正；MRL 已验收为负，十折回归待闭环", 2),
            ("ncRNA 保留 accuracy / macro-F1 / balanced accuracy；MRL 与 TE-EL 保留 PCC。",
             "MRL full_v3：R² Δ=-0.468847（不显著为正）；图 8B 待 TE/EL 十折，其余任务继续验收。", 1),
            ("合并原第 25、28a、36 页信息；splice 数据留档不进主文。新增任务必须完成预注册与统计后再加入。",
             "working 面板保留全部负结果；旧 177 类只作敏感性，不替代官方 88 类；图 8B 暂不入片。", 2),
            ("ROUND26 / ROUND28 / T-A129；frozen backbone 双头对照",
             "round33/34 full results + external QA；分类与回归分图", 2),
        ],
        15: [
            ("NucleicBERT 8/8 正增益；ncRNA 三指标提升",
             "NB 主 benchmark 8/8 正迁移；原论文下游任务未显著为正", 1),
            ("MRL/TE-EL 增益较小；新任务未完成前不得扩写为普遍优势",
             "NB 其余官方资产 blocked；MRL 已验收为负，TE/EL 十折未闭环", 1),
        ],
        19: [
            ("三组 MRL 已预注册；5-seed、early-stopping、paired t-test+Holm。",
             "MRL 与 HEK fold0–2 TE/EL 已 QA，主指标均不显著为正；fold3 起受门推进。", 1),
        ],
        20: [
            ("大样本 ECE 改善；NB 8/8", "大样本 ECE 改善；NB 主 benchmark 8/8", 1),
            ("ncRNA 三指标提升", "原论文任务未显著为正", 1),
            ("新增 MRL 未过显著门；不外推", "官方资产与全协议未闭环；不外推", 1),
        ],
    }
    for slide_number, items in replacements.items():
        slide = prs.slides[slide_number - 1]
        for old, new, expected in items:
            replace_run_text(slide, old, new, expected - 1 if expected == 2 else expected)
            if expected == 2:
                replace_run_text(slide.notes_slide, old, new, 1)


def update_fold3_te_status(prs):
    replace_run_text(
        prs.slides[18],
        "MRL 与 HEK fold0–2 TE/EL 已 QA，主指标均不显著为正；fold3 起受门推进。",
        "MRL 与 HEK fold0–2 TE/EL、fold3 TE 已 QA，均不显著为正；fold3 EL 待续接。",
        1,
    )


def rebuild_fig6(slide):
    for index, shape in enumerate(list(slide.shapes)):
        if index < 6:
            continue
        if hasattr(shape, "text") and shape.text.startswith("与后文关系"):
            continue
        remove(shape)

    # Three coordinated panels with one continuous left-to-right data flow.
    panels = [(0.55, 3.72, "paleBlue"), (4.52, 4.28, "surface"), (9.05, 3.73, "paleTeal")]
    for x, w, fill in panels:
        box(slide, x, 1.55, w, 4.96, fill, "rule")
    text(slide, "A  FROZEN SEQUENCE REPRESENTATION", 0.76, 1.73, 3.30, 0.32, 10.3, "blue", True)
    text(slide, "B  PAIRWISE 2D DECISION", 4.76, 1.73, 3.80, 0.32, 10.3, "blue", True)
    text(slide, "C  CONSTRAINED DECODING + CALIBRATION", 9.25, 1.73, 3.32, 0.32, 9.2, "teal", True)

    # A: sequence, frozen encoder and embedding tensor.
    text(slide, "RNA sequence · length L", 0.82, 2.06, 2.90, 0.22, 8.6, "muted")
    for index, token in enumerate(["A", "U", "G", "C", "…", "U"]):
        labeled_box(slide, token, 0.78 + index * 0.48, 2.38, 0.38, 0.38,
                    "white", "rule", 9.2, "ink", True)
    connector(slide, 2.41, 2.76, 2.41, 3.20, "blue")
    box(slide, 0.92, 3.20, 3.02, 0.92, "blue", "blue", True, 0.6)
    text(slide, "RiNALMo-giga · 650M", 1.08, 3.30, 2.70, 0.28, 12, "white", True, PP_ALIGN.CENTER)
    text(slide, "33 Transformer blocks · 1280-d output", 1.08, 3.62, 2.70, 0.24, 8.8, "white", False, PP_ALIGN.CENTER)
    text(slide, "frozen backbone · 0 trainable parameters", 1.08, 3.87, 2.70, 0.24, 8.8, "white", False, PP_ALIGN.CENTER)
    connector(slide, 2.41, 4.12, 2.41, 4.46, "blue")
    matrix(slide, 1.32, 4.46, 0.18, 7, ("paleBlue", "blue2", "teal"))
    text(slide, "h: L × 1280", 2.76, 4.62, 1.34, 0.34, 9.6, "blue", True, PP_ALIGN.CENTER)
    text(slide, "one embedding per nucleotide", 2.76, 5.00, 1.34, 0.42, 8.6, "muted", False, PP_ALIGN.CENTER)
    labeled_box(slide, "cached or recomputed on raw sequences", 0.91, 5.73, 3.02, 0.42,
                "white", "rule", 8.6, "muted")
    text(slide, "Output: contextual sequence features", 0.91, 6.20, 3.02, 0.22, 8.6, "blue", True, PP_ALIGN.CENTER)

    # B: symmetric pair representation, pair tensor and 2D scorer.
    labeled_box(slide, "h_i", 4.82, 2.22, 0.62, 0.38, "white", "rule", 10, "blue", True)
    labeled_box(slide, "h_j", 5.55, 2.22, 0.62, 0.38, "white", "rule", 10, "blue", True)
    text(slide, "sum  h_i + h_j\nproduct  h_i ⊙ h_j\ndistance  |h_i − h_j|", 6.36, 2.16, 2.10, 0.66,
         8.6, "ink", False, PP_ALIGN.LEFT)
    text(slide, "zᵢⱼ = W_z [sum ; product ; distance]\n+ relative-offset(i,j)",
         4.82, 2.84, 3.68, 0.52, 8.6, "muted", False, PP_ALIGN.CENTER)
    matrix(slide, 4.95, 3.45, 0.18, 6, ("paleBlue", "blue2", "teal"))
    text(slide, "L × L × 128", 4.86, 4.62, 1.34, 0.24, 8.6, "muted", False, PP_ALIGN.CENTER)
    connector(slide, 6.03, 3.96, 6.60, 3.96, "blue")
    labeled_box(slide, "ResNet2D scorer\n1×1 stem → 4 residual blocks (3×3)\n→ 1×1 head",
                6.60, 3.35, 1.92, 1.18, "blue2", "blue2", 8.6, "white", True)
    labeled_box(slide, "symmetric L × L pair logits\nsingle forward · no autoregression",
                4.84, 5.03, 3.66, 0.60, "white", "rule", 8.8, "blue", True)
    labeled_box(slide, "NLL + KL distillation + RLCD + soft-ECE",
                4.84, 5.78, 3.66, 0.38, "paleGold", "paleGold", 8.6, "ink", False)
    text(slide, "decision head: 812,375 trainable parameters", 4.84, 6.22, 3.66, 0.22,
         8.8, "blue", True, PP_ALIGN.CENTER)

    # C: prior injection, constrained DP and two output branches.
    labeled_box(slide, "scoreᵢⱼ + w_turner × TurnerPriorᵢⱼ\nvalid-pair mask: canonical pairs + min-loop",
                9.34, 2.20, 3.16, 0.68, "white", "rule", 8.6, "ink", True)
    connector(slide, 10.92, 2.88, 10.92, 3.24, "teal")
    labeled_box(slide, "non-crossing max-product DP\nC-batched exact decode · legal by construction",
                9.34, 3.24, 3.16, 0.86, "teal", "teal", 9.0, "white", True)
    text(slide, "O(L³) exact decode; no partition-function computation",
         9.36, 4.17, 3.12, 0.30, 8.6, "muted", False, PP_ALIGN.CENTER)
    labeled_box(slide, "STRUCTURE\ndot-bracket + legal base pairs",
                9.34, 4.78, 1.50, 0.88, "white", "rule", 8.5, "teal", True)
    labeled_box(slide, "CONFIDENCE\nsigmoid → VL0 affine → SC / ECE",
                11.00, 4.78, 1.50, 0.88, "white", "rule", 8.5, "blue", True)
    text(slide, "structure branch", 9.36, 5.76, 1.46, 0.24, 8.5, "muted", False, PP_ALIGN.CENTER)
    text(slide, "calibration branch", 11.00, 5.76, 1.50, 0.24, 8.5, "muted", False, PP_ALIGN.CENTER)
    labeled_box(slide, "legal RNA secondary structure + per-structure confidence",
                9.34, 6.06, 3.16, 0.42, "paleTeal", "paleTeal", 8.6, "teal", True)

    # Cross-panel tensor hand-offs, drawn above the panels.
    connector(slide, 4.28, 3.82, 4.50, 3.82, "blue", 1.4)
    text(slide, "h", 4.28, 3.55, 0.22, 0.22, 8.5, "blue", True, PP_ALIGN.CENTER)
    connector(slide, 8.80, 3.82, 9.03, 3.82, "teal", 1.4)
    text(slide, "logits", 8.69, 3.55, 0.45, 0.22, 8.5, "teal", True, PP_ALIGN.CENTER)


def final_asset(name):
    value = os.environ.get(name)
    if not value:
        return None
    path = Path(value)
    assert path.is_file() and path.stat().st_size > 100_000, (name, path)
    return path


def replace_slide_and_notes(prs, slide_number, old, new):
    replace_run_text(prs.slides[slide_number - 1], old, new, 1)
    replace_run_text(prs.slides[slide_number - 1].notes_slide, old, new, 1)


def clone_regression_slide(prs, image_path):
    source = prs.slides[11]
    picture = next(shape for shape in source.shapes if shape.shape_type == 13)
    new_slide = prs.slides.add_slide(source.slide_layout)
    for shape in list(new_slide.shapes):
        remove(shape)
    for shape in source.shapes:
        if shape.shape_type != 13:
            new_slide.shapes._spTree.insert_element_before(deepcopy(shape._element), "p:extLst")
    new_slide.shapes.add_picture(str(image_path), picture.left, picture.top, picture.width, picture.height)
    slide_ids = prs.slides._sldIdLst
    new_id = slide_ids[-1]
    slide_ids.remove(new_id)
    slide_ids.insert(12, new_id)
    return new_slide


def renumber_after_regression_insert(prs):
    replace_run_text(prs.slides[1], "11–14", "11–15", 1)
    replace_run_text(prs.slides[1], "15–21", "16–22", 1)
    for index, slide in enumerate(prs.slides):
        if index < 12:
            continue
        matches = []
        for shape in slide.shapes:
            if (getattr(shape, "has_text_frame", False) and shape.left > prs.slide_width * 0.85
                    and shape.top > prs.slide_height * 0.85 and shape.text.strip().isdigit()):
                matches.append(shape)
        if not matches and index == len(prs.slides) - 1:
            assert any(getattr(shape, "has_text_frame", False) and shape.text.strip() == "END"
                       for shape in slide.shapes)
            continue
        assert len(matches) == 1, (index + 1, len(matches))
        runs = [run for paragraph in matches[0].text_frame.paragraphs for run in paragraph.runs if run.text.strip()]
        assert len(runs) == 1
        runs[0].text = f"{index + 1:02d}"


def apply_final_assets(prs):
    fig05 = final_asset("FIG05_FINAL")
    fig07 = final_asset("FIG07_FINAL")
    fig08a = final_asset("FIG08A_FINAL")
    fig08b = final_asset("FIG08B_FINAL")
    if fig05:
        replace_picture(prs.slides[8], fig05)
        replace_slide_and_notes(prs, 9, "主图 5｜历史 A100 帕累托集；4090 共同协议重测中",
                                "主图 5｜RTX 4090：精度-吞吐与共同 batch 网格")
        replace_run_text(prs.slides[8], "当前图仅作 A100 历史参照；正式版使用同一 4090、TS0、raw-sequence 端到端 B=1/2/4 结果。",
                         "同一 RTX 4090、TS0、raw-sequence 端到端；A 为 B=1 精度-吞吐，B 为共同 B=1/2/4 网格。", 1)
        replace_slide_and_notes(prs, 9,
                                "旧 124.78 为 cached-embedding/head-only，不能混入 raw-sequence 前沿；当前仅四个系统 B=1 已通过，完整网格待验收。",
                                "六系统均含自模型；预测跨 batch 一致，UFold 官方权重不可得，不混入 A100 或 head-only 数据。")
        replace_slide_and_notes(prs, 9, "A100 baseline board + §15.79；Pareto 程序逐点判定",
                                "speed_4090_round33/34 full JSON；共同 4090 协议与独立复核")
        replace_run_text(prs.slides[14], "4090 raw-seq B=1：r2d 22.28 seq/s；旧 head-only 峰值 124.78",
                         "4090 raw-seq 全网格：六系统 B=1/2/4；含 r2d/plana/xens2", 1)
        replace_run_text(prs.slides[14], "共同 batch 网格未闭环；禁止跨口径拼接速度或 xens2 F1",
                         "UFold 权重 blocked；不混用 A100、cached embedding 或跨协议结果", 1)
        replace_run_text(prs.slides[19], "raw-seq B=1：r2d 22.28 seq/s", "4090 raw-seq B=1/2/4 全网格", 1)
        replace_run_text(prs.slides[19], "旧 head-only 峰值 124.78；共同 batch 网格待验收",
                         "预测跨 batch 一致；UFold blocked", 1)
    if fig07:
        replace_picture(prs.slides[10], fig07)
        replace_slide_and_notes(prs, 11, "主图 7｜主 benchmark 8/8 正迁移；原论文三任务不显著为正",
                                "主图 7｜主 benchmark 与原论文任务：五种子完整对照")
        replace_run_text(prs.slides[10], "主 benchmark 与原论文任务分开展示；acceptor/donor/CPEB3 全量五种子均不显著为正。",
                         "主 benchmark、二级结构、acceptor/donor 与 CPEB3 分区；全部正负结果及 Holm 保留。", 1)
        replace_slide_and_notes(prs, 11,
                                "8/8 仅指主 benchmark；secondary-structure 数据门已过、五种子待验收；contact-distance/shuffle 保持 blocked，不作全覆盖声明。",
                                "二级结构五种子已入图；contact-distance/shuffle 仍 blocked，不作全覆盖或去污染泛化声明。")
        replace_slide_and_notes(prs, 11, "ROUND26/T-A99 + round34 downstream QA；其余原论文任务待官方资产",
                                "ROUND26/T-A99 + round34/35 full results + external QA")
        replace_run_text(prs.slides[19], "原论文任务未显著为正", "原论文任务按完整证据呈现", 1)
    if fig08a:
        replace_picture(prs.slides[11], fig08a)
        replace_slide_and_notes(prs, 12, "主图 8A｜官方分类未显著为正；MRL 已验收为负，十折回归待闭环",
                                "主图 8A｜分类任务：结构、Spliceator 与官方 88 类 ncRNA")
        replace_run_text(prs.slides[11], "MRL full_v3：R² Δ=-0.468847（不显著为正）；图 8B 待 TE/EL 十折，其余任务继续验收。",
                         "十项结构任务与两项分类家族分置；五种子配对、任务内 Holm，全部正负结果保留。", 1)
        replace_slide_and_notes(prs, 12, "working 面板保留全部负结果；旧 177 类只作敏感性，不替代官方 88 类；图 8B 暂不入片。",
                                "分类结果已闭环；MRL 与 TE/EL 回归见下一页图 8B。")
        replace_slide_and_notes(prs, 12, "round33/34 full results + external QA；分类与回归分图",
                                "round33/34 full results + structure aggregate + external QA")
    if fig08b:
        assert fig08a, "Figure 8B requires the final Figure 8A slide"
        new_slide = clone_regression_slide(prs, fig08b)
        replace_run_text(new_slide, "主图 8A｜分类任务：结构、Spliceator 与官方 88 类 ncRNA",
                         "主图 8B｜回归任务：official MRL 与三细胞 TE/EL", 1)
        replace_run_text(new_slide, "十项结构任务与两项分类家族分置；五种子配对、任务内 Holm，全部正负结果保留。",
                         "MRL 使用 R²；TE/EL 先在每 seed 内十折平均 Spearman，再做五 seed 配对。", 1)
        replace_run_text(new_slide, "分类结果已闭环；MRL 与 TE/EL 回归见下一页图 8B。",
                         "七项回归结果全部入图；所有负结果保留，分类结果见上一页图 8A。", 1)
        replace_run_text(new_slide, "round33/34 full results + structure aggregate + external QA",
                         "MRL full_v3 + six ten-fold CV aggregates + external QA", 1)
        renumber_after_regression_insert(prs)
        replace_run_text(prs.slides[19], "MRL 与 HEK fold0–2 TE/EL、fold3 TE 已 QA，均不显著为正；fold3 EL 待续接。",
                         "MRL、TE/EL 十折与结构任务已聚合；全部正负结果保留。", 1)
    if fig07 or fig08a or fig08b:
        conclusion_slide = prs.slides[15 if fig08b else 14]
        logic_slide = prs.slides[20 if fig08b else 19]
        replace_run_text(conclusion_slide, "NB 主 benchmark 8/8 正迁移；原论文下游任务未显著为正",
                         "NB 主 benchmark 8/8 正迁移；原论文任务按五种子证据报告", 1)
        replace_run_text(conclusion_slide, "NB 二级结构待五种子；contact/distance/shuffle blocked；MRL 为负，TE/EL 未闭环",
                         "contact/distance/shuffle blocked；已测正负结果完整保留，未覆盖任务不外推", 1)
        replace_run_text(logic_slide, "官方资产与全协议未闭环；不外推",
                         "已测结果完整保留；未覆盖任务不外推", 1)
    return {"fig05": bool(fig05), "fig07": bool(fig07), "fig08a": bool(fig08a), "fig08b": bool(fig08b)}


def main():
    prs = Presentation(SRC)
    replace_picture(prs.slides[4], FIG / "fig02_speed_benchmark_3840x1920.png")
    replace_picture(prs.slides[5], FIG / "fig03a_ablation_matrix_3840x1920.png")
    replace_picture(prs.slides[6], FIG / "fig03b_ood_signature_3840x1920.png")
    replace_picture(prs.slides[7], FIG / "fig04_ece_confidence_3840x1920.png")
    replace_picture(prs.slides[8], FIG / "fig05_accuracy_speed_pareto_3840x1920.png")
    replace_picture(prs.slides[10], FIG / "fig07_nucleicbert_combined_working_3840x1920.png")
    replace_picture(prs.slides[11], FIG / "fig08a_rinalmo_classification_working_3840x1920.png")
    for shape in list(prs.slides[5].shapes):
        if hasattr(shape, "text") and shape.text.startswith("那么这部分我有一些建议"):
            remove(shape)
    for slide_index, prefix in ((8, "更新了，到 4090"), (11, "Splice 的没放")):
        for shape in list(prs.slides[slide_index].shapes):
            if hasattr(shape, "text") and shape.text.startswith(prefix):
                remove(shape)
    rebuild_fig6(prs.slides[9])
    update_speed_provenance(prs)
    update_downstream_provenance(prs)
    update_secondary_data_status(prs)
    update_fold3_te_status(prs)
    update_version_labels(prs)
    applied = apply_final_assets(prs)
    prs.save(OUT)
    print(f"ROUND34_DECK_WRITTEN {OUT} slides={len(prs.slides)} final_assets={applied}")


if __name__ == "__main__":
    main()
