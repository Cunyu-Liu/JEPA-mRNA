#!/usr/bin/env python3
"""Add slide 26 to RNA-JEV_项目进展_20260924.pptx: the 10-19 round
(v3.20->v3.26, TS2 axes exhausted, 124.78 seq/s, confidence board,
submission PDF). Mirrors slide 25's layout/typography (TextBox family,
title teal 0E7C86, body dark). Numbers pre-checked: checkers 8/8+9/9,
PDF 8/8 content checks, figures p23 verified."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import copy

P = "/Users/liucunyu/Documents/课题idea 数据 结论/浙江大学/RNA 相关模型/RNA-JEV/ppt/RNA-JEV_项目进展_20260924.pptx"
p = Presentation(P)
src = p.slides[-1]  # slide 25 template

lay = src.slide_layout
new = p.slides.add_slide(lay)

TITLE = "预印本投稿包就绪 · v3.26"
SUB = "TS2 集成轴双向穷尽 + 速度 124.8 seq/s + 置信度全板 + 24 页 PDF 投稿包（全链三遍核查）"
BODY = ("本轮闭环：① 板协议 8/8 在 4090 双硬件逐位复现（max |Δ|=0.0001）——每个数字都有两条硬件证据；"
        "② TS2 攻关七战线收官：集成轴双向穷尽（r2d 成员侧 + plana 桶 2/3/4-seed 0.8604/0.8581/0.8526 单调递减），"
        "0.8604 终判如实写入范式边界；③ 速度线：B=256 fp16 = 124.78 seq/s（F1 逐位无损），"
        "RiNALMo-ft 同节点复现 14.5 seq/s → 8.6× 同 GPU 对比；④ Jev 式结构级置信度 8-split 全板（大样本 4/4 校准迁移，"
        "小样本族敏感如实报告）；⑤ 2025-26 参照补全（HydraRNA/RNADiffFold）；"
        "⑥ 投稿包：v3.26 + Data/Code availability + 附录 D 图注 → 24 页 PDF（两幅主图 p23，8/8 内容检查）。")
FOOT = "RNA-JEPA 决策模型线 · 2026-10-19 · 提交 333ee32"
PG = "26"

def add_tb(name, text, pos, size, color, bold=False, w=9825280, h=400000):
    tb = new.shapes.add_textbox(pos[0], pos[1], w, h)
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    # 1.5x line spacing per user requirement
    from pptx.oxml.ns import qn
    pPr = tf.paragraphs[0]._pPr if tf.paragraphs[0]._pPr is not None else None
    para = tf.paragraphs[0]
    para.line_spacing = 1.5
    run = para.add_run()
    run.text = text
    f = run.font
    f.name = "微软雅黑"
    f.size = Pt(size)
    f.color.rgb = RGBColor.from_string(color)
    f.bold = bold
    return tb

# geometry copied from slide 25 shapes
add_tb("T1", TITLE, (493776, 310896), 15, "0E7C86", bold=True, h=300000)
add_tb("T2", SUB, (493776, 548640), 16, "0F1B2D", bold=True, h=500000)

# body block (single flowing block, matches TextBox 8 style; 1.5x spacing)
tb = new.shapes.add_textbox(493776, 1188720, 9825280, 4000000)
tb.name = "T3"
tf = tb.text_frame
tf.word_wrap = True
para = tf.paragraphs[0]
para.line_spacing = 1.5
run = para.add_run()
run.text = BODY
f = run.font
f.name = "微软雅黑"; f.size = Pt(11); f.color.rgb = RGBColor.from_string("333A45")

add_tb("T4", "预印本状态：v3.26 投稿就绪（24 页 PDF + markdown + standalone HTML 四地一致）。\n下一步：导师审阅 → bioRxiv/arXiv 投稿；可选 T-A84 消融重训（6 个非板路径历史 ckpt）。",
       (493776, 5394960), 11, "0F1B2D", h=800000)
add_tb("T5", FOOT, (493776, 6419088), 10, "8A93A0", h=280000)
add_tb("T6", PG, (8769096, 6419088), 11, "1F3A5F", h=280000, w=500000)

p.save(P)
print("slide 26 added")
