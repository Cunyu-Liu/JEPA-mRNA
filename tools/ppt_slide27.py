#!/usr/bin/env python3
"""Slide 27: the paper's three innovation points (user-specified framing):
1) main-benchmark SOTA 2) speed-benchmark SOTA 3) confidence-score
correlation with downstream task scores. Plus engineering highlights
(cleaning/splitting/leak-audit/baseline reproduction) as a supporting band.
Follows formal-work-report recipe: claim band, evidence field (3 innovation
columns + engineering band), implication, footer. 1.5x line spacing,
Microsoft YaHei throughout."""
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor

P = "/Users/liucunyu/Documents/课题idea 数据 结论/浙江大学/RNA 相关模型/RNA-JEV/ppt/RNA-JEV_项目进展_20260924.pptx"
p = Presentation(P)
lay = p.slides[-1].slide_layout
new = p.slides.add_slide(lay)

TEAL = "0E7C86"; DARK = "0F1B2D"; BODY = "333A45"; GREY = "8A93A0"
NAVY = "1F3A5F"

def tb(name, x, y, w, h, text, size, color, bold=False, align=None):
    box = new.shapes.add_textbox(x, y, w, h)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    para = tf.paragraphs[0]
    para.line_spacing = 1.5
    if align:
        para.alignment = align
    r = para.add_run(); r.text = text
    f = r.font
    f.name = "微软雅黑"; f.size = Pt(size)
    f.color.rgb = RGBColor.from_string(color); f.bold = bold
    return box

L = 493776; W = 9825280
# claim band
tb("T1", L, 310896, W, 320000, "论文三个创新点 · 工程亮点（v3.26 定稿口径）", 15, TEAL, bold=True)
tb("T2", L, 548640, W, 500000,
   "主 benchmark SOTA × 速度 benchmark SOTA × 置信度-下游分数相关性——三个创新点均已实测闭环", 16, DARK, bold=True)

# three innovation columns (evidence field), each 1/3 width
CW = 3170000; GAP = 155280
Y0 = 1250000; CH = 3100000
inno = [
    ("创新点 1 · 主 benchmark",
     "7/8 split 超 SOTA（唯一真负 TS2 写成范式边界）\n"
     "· xens2 板协议 8/8 双硬件逐位复现\n"
     "· 对干净参照 6/8 单模型胜、7/8 集成胜\n"
     "· 参照去污染翻案（TS3/hard 胜）"),
    ("创新点 2 · 速度 benchmark",
     "124.78 seq/s（B=256 fp16，F1 逐位无损）\n"
     "· 同节点 RiNALMo-ft 复现 14.5 seq/s → 8.6×\n"
     "· 对 UFold/MXfold2/RNAformer 跨硬件 2-5×\n"
     "· 批轴见顶：B=64→256 阶梯 96→125"),
    ("创新点 3 · 置信度",
     "结构级置信度与下游分数强相关\n"
     "· Spearman 0.709（TS0，n=1281，p<1e-195）\n"
     "· ECE 0.0315 vs 原始 0.1604\n"
     "· 8-split 全板：大样本 4/4 迁移"),
]
for i, (head, body) in enumerate(inno):
    x = L + i * (CW + GAP)
    tb(f"H{i}", x, Y0, CW, 400000, head, 13, NAVY, bold=True)
    tb(f"B{i}", x, Y0 + 430000, CW, CH - 430000, body, 10.5, BODY)

# engineering highlights band (implication band)
tb("E0", L, 4500000, W, 330000, "工程性亮点（评测卫生 = 论文可信度基座）", 12.5, TEAL, bold=True)
tb("E1", L, 4860000, W, 1100000,
   "① 数据集清洗：tr1c 去污染语料（两次自我撤回驱动）；② 测评集切割：9-split 冻结 + SHA-256 清单；"
   "③ 数据泄露检测：全矩阵审计（TS0 1087 重叠撤回 / TS3 参照 79% 污染翻案）；"
   "④ 其他模型完美复现：RiNALMo-ft 双 split 逐位 0.0001 + RNAformer/UFold 官方权重同口径重测。", 10.5, BODY)

# footer
tb("F1", L, 6419088, 8000000, 280000, "RNA-JEPA 决策模型线 · 2026-10-19 · 提交 333ee32", 10, GREY)
tb("F2", 8769096, 6419088, 500000, 280000, "27", 11, NAVY)

p.save(P)
print("slide 27 added")
