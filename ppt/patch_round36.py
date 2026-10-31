#!/usr/bin/env python3
import argparse
import json
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

FONT = "微软雅黑"
SPLITS = ["TS0", "new", "TS1", "TS2", "TS3", "hard", "ArchII", "TSB"]


def replace_text(slide, old, new):
    matches = []
    for shape in slide.shapes:
        frames = []
        if shape.has_text_frame:
            frames.append(shape.text_frame)
        if shape.has_table:
            frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
        for frame in frames:
            for paragraph in frame.paragraphs:
                paragraph_matched = False
                for run in paragraph.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        matches.append(run)
                        paragraph_matched = True
                if paragraph_matched:
                    continue
                full = "".join(run.text for run in paragraph.runs)
                if old in full:
                    paragraph.runs[0].text = full.replace(old, new)
                    for run in paragraph.runs[1:]:
                        run.text = ""
                    matches.append(paragraph.runs[0])
    if not matches:
        raise AssertionError(f"text not found: {old}")
    return matches


def set_shape_text(shape, text, size=None, bold=None):
    shape.text_frame.clear()
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold


def clone_slide(prs, source_index, insert_index):
    source = prs.slides[source_index]
    target = prs.slides.add_slide(prs.slide_layouts[0])
    for shape in source.shapes:
        target.shapes._spTree.insert_element_before(deepcopy(shape.element), "p:extLst")
    slide_id = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(slide_id)
    prs.slides._sldIdLst.insert(insert_index, slide_id)
    return prs.slides[insert_index]


def set_notes(slide, text):
    frame = slide.notes_slide.notes_text_frame
    if frame is not None:
        frame.text = text


def patch_speed_pages(prs):
    slide = prs.slides[4]
    replace_text(slide, "主图 2｜历史 head-only 峰值 124.78 seq/s，禁止跨口径比较",
                 "主图 2｜124.78 是 head-only 峰值，不是 4090 复刻失败")
    replace_text(slide,
                 "124.78 属于 r2dtr1c-s0 的历史 head-only 峰值；当前 raw-sequence B=1 为 22.28 seq/s（F1=0.668181），二者不得混用。",
                 "同 B=1：head-only 21.45、raw-sequence 22.28 seq/s，二者接近；124.78 主要来自 B=256+fp16 饱和加速及不同计时边界，并非 4090 变慢。当前未证明速度 SOTA。")
    set_notes(slide,
              "同 B=1 时，head-only=21.45、raw-sequence=22.28 seq/s，二者接近。124.78 主要来自 B=256+fp16 饱和加速和不同计时边界，不是 4090 复刻失败。当前只有已测五系统内部 Pareto，未证明全模型速度 SOTA。")

    slide = prs.slides[8]
    replace_text(slide,
                 "RNAformer 因 1024-nt padding 未能在独占窗口完成 B=1，保持 blocked；UFold 官方权重不可得；旧 124.78 seq/s 仅历史值，禁止混入 raw-sequence 前沿。",
                 "当前仅为已测五系统内部 Pareto，不是全模型速度 SOTA；RNAformer/UFold 同机完整点补齐后再判定。旧 124.78 seq/s 为 head-only 历史值，禁止混入 raw-sequence 前沿。")
    set_notes(slide,
              "RTX 4090、TS0 n=1288、raw-sequence 端到端。当前 15/18 格只支持已测五系统内部比较，不能称为全模型速度 SOTA。RNAformer/UFold 同机完整点仍在补齐，旧 124.78 只作 head-only 历史参照。")

    slide = prs.slides[15]
    replace_text(slide,
                 "RTX 4090 raw-seq 15 格：r2d 22.28→37.42 seq/s（B=1→4）；xens2 7.47→12.38；旧 head-only 124.78 仅历史",
                 "4090 raw-seq：r2d 22.28→37.42、xens2 7.47→12.38 seq/s（B=1→4）；速度优化中，当前未证明 SOTA")
    replace_text(slide,
                 "RNAformer 1024-nt padding 未能完成 B=1，UFold 权重不可得；跨 batch 仅数值等价",
                 "RNAformer/UFold 同机完整点未闭环；历史 124.78 为 head-only；跨 batch 仅数值等价")

    slide = prs.slides[20]
    replace_text(slide,
                 "raw-seq 与共同 batch 网格均已完成（15/18 格）；旧 head-only 124.78 仅历史值",
                 "15/18 格只支持已测五系统内部比较；旧 head-only 124.78 仅历史值，当前未证明速度 SOTA")


def patch_protocol_page(prs):
    slide = prs.slides[17]
    for shape in list(slide.shapes):
        if shape.has_text_frame and shape.text.startswith("没看懂"):
            shape.element.getparent().remove(shape.element)
    replacements = {
        "模型、硬件、计时范围必须同时写": ("三种速度口径：输入、计时边界与当前判定", 20.5, True),
        "模型身份": ("head-only｜缓存表征", 11.5, True),
        "124.78 属于 r2dtr1c-s0 的历史 head-only 峰值，不是 xens2；F1=0.667971。": (
            "cached embedding → 812,375 head → C 解码；不含编码器。124.78 为 B=256/fp16 峰值。", 9.5, False),
        "计时范围": ("forward-only｜仅前向", 11.5, True),
        "含 cached embedding、812,375 head 与 C 解码；不包含 raw-sequence 编码器前向。": (
            "raw sequence → 编码器 + 任务头 → 分数/概率；不含 DP。RiNALMo-ft 14.5 seq/s 属此类。", 9.5, False),
        "硬件边界": ("raw sequence｜端到端", 11.5, True),
        "UFold/RNAformer/MXfold2 的历史 head-only 点为 A100 板；4090 raw-sequence 点不与其混算 Pareto。": (
            "raw sequence → 编码器 + 任务头 + 必要 DP；r2d B=1 为 22.28 seq/s，三步全部计入。", 9.5, False),
        "参照限制": ("为何差这么多", 11.5, True),
        "RiNALMo-ft 14.5 seq/s 为 forward-only，单列，不能冒充完整系统延迟。": (
            "B=1 两者已接近（21.45 vs 22.28）；124.78 主要来自 B=256/fp16 饱和加速与计时边界，并非硬件退化。", 9.5, False),
        "有效性": ("当前判定", 11.5, True),
        "tp/fp/fn=24406/13778/10485；批大小不改变预测，B=384/512 只作饱和检查。": (
            "已测五系统内部 Pareto ≠ 全模型速度 SOTA；同机基线未闭环，后续集中做无损推理优化。", 9.5, False),
    }
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text in replacements:
            text, size, bold = replacements[shape.text]
            set_shape_text(shape, text, size=size, bold=bold)
    replace_text(slide, "本页为历史 head-only 阶梯；不得与 raw-sequence 或 forward-only 读数混用。",
                 "三种计时边界见右侧；124.78 与 22.28 不是同协议，不能称为 4090 复刻失败。")
    set_notes(slide,
              "head-only：cached embedding→决策头→必要 DP，不含编码器；forward-only：raw sequence→编码器+任务头，不含 DP；raw-sequence：包含编码器、任务头和必要 DP。B=1 两者已接近，124.78 主要来自 B=256+fp16 饱和加速。当前未证明速度 SOTA。")


def patch_leakage_page(prs, source):
    new_slide = clone_slide(prs, 16, 17)
    replace_text(new_slide, "附录 A / 全板", "附录 A2 / 泄露处理")
    replace_text(new_slide, "全板敏感性：xens2=5/8，最佳单 checkpoint=4/8",
                 "泄露处理全板：污染格剔除，TR0 系 TestSetB 统一 clean-181")
    replace_text(new_slide,
                 "完整保留基线；统一 xens2 与历史 r2d_tr1 集成单列，避免旧表混用。",
                 "19 系统 × 8 split，与附录 A 原始全板对应；×=污染剔除，—=缺测/未复算，均不按 0 分。")
    replace_text(new_slide, "inter-fam/biophys 为官方 checkpoint 项目 GT 复测；原文 Table 4 见主图 1 注释",
                 "strict micro-F1 + 项目 canonical-nested GT；数值来自原始 result/audit JSON")
    page_shape = next(shape for shape in new_slide.shapes if shape.has_text_frame and shape.text == "17")
    page_shape.text_frame.paragraphs[0].runs[0].text = "18"
    table = next(shape.table for shape in new_slide.shapes if shape.has_table)
    assert len(table.rows) == len(source["rows"]) + 1
    for row_index, row in enumerate(source["rows"], start=1):
        for col_index, split in enumerate(SPLITS, start=1):
            target = table.cell(row_index, col_index)
            runs = target.text_frame.paragraphs[0].runs
            if runs:
                runs[0].text = row["cells"][split]["display"]
            else:
                target.text = row["cells"][split]["display"]
    for row in table.rows:
        for target in row.cells:
            for paragraph in target.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(8)
    note = next(shape for shape in new_slide.shapes if shape.has_text_frame and shape.text.startswith("RNAformer-if/biophys"))
    set_shape_text(note,
                   "†=TestSetB clean-181（剔除 247 条 TR0 重叠）；×=精确重叠污染格剔除；—=缺测或未复算。"
                   "‡=特殊 checkpoint/审计边界；RNAformer-if TS2=0.9043 为项目 GT 复测，原文 0.802 属全文配对口径，不能互换。",
                   size=8.5, bold=False)
    return new_slide


def patch_page_numbers(prs):
    for index in range(18, len(prs.slides)):
        slide = prs.slides[index]
        expected = str(index + 1)
        candidates = [shape for shape in slide.shapes if shape.has_text_frame and shape.text.isdigit()]
        if candidates:
            runs = candidates[-1].text_frame.paragraphs[0].runs
            if runs:
                runs[0].text = expected
            else:
                candidates[-1].text = expected
    replace_text(prs.slides[1], "主文 10 图 + 4 个附录 + 1 页论文逻辑链",
                 "主文 10 图 + 5 个附录 + 1 页论文逻辑链")
    replace_text(prs.slides[1], "16–22", "16–23")
    replace_text(prs.slides[-1], "主文 10 图 · 附录 4 组 · 论文逻辑链 1 页 · 微软雅黑 · 1.5 倍行距",
                 "主文 10 图 · 附录 5 组 · 论文逻辑链 1 页 · 微软雅黑 · 1.5 倍行距")


def normalize_font_and_spacing(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if shape.has_table:
                frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
            for frame in frames:
                for paragraph in frame.paragraphs:
                    paragraph.line_spacing = 1.5
                    for run in paragraph.runs:
                        run.font.name = FONT


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--deck", type=Path, required=True)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    source = json.loads(args.source.read_text(encoding="utf-8"))
    prs = Presentation(args.deck)
    assert len(prs.slides) == 22
    replace_text(prs.slides[0], "2026-10-31 · Round 35", "2026-10-31 · Round 36")
    replace_text(prs.slides[15], "结论口径：Round 35 统一复核", "结论口径：Round 36 统一复核")
    replace_text(prs.slides[-1], "RNA-JEV / ROUND 35", "RNA-JEV / ROUND 36")
    patch_speed_pages(prs)
    patch_protocol_page(prs)
    patch_leakage_page(prs, source)
    patch_page_numbers(prs)
    normalize_font_and_spacing(prs)
    assert len(prs.slides) == 23
    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)
    print(f"ROUND36_DECK_WRITTEN {args.out}")


if __name__ == "__main__":
    main()
